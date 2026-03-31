"""Generate the tpacmon technical report as DOCX."""

from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
import datetime


def setup_styles(doc):
    """Configure document styles."""
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Calibri"
    font.size = Pt(11)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.15

    for level in range(1, 5):
        heading = doc.styles[f"Heading {level}"]
        heading.font.name = "Calibri"
        heading.font.color.rgb = RGBColor(0x1A, 0x3C, 0x6E)
        if level == 1:
            heading.font.size = Pt(18)
            heading.paragraph_format.space_before = Pt(24)
        elif level == 2:
            heading.font.size = Pt(14)
            heading.paragraph_format.space_before = Pt(18)
        elif level == 3:
            heading.font.size = Pt(12)
            heading.paragraph_format.space_before = Pt(12)

    # Code style
    code_style = doc.styles.add_style("Code", WD_STYLE_TYPE.PARAGRAPH)
    code_style.font.name = "Consolas"
    code_style.font.size = Pt(9)
    code_style.paragraph_format.space_before = Pt(3)
    code_style.paragraph_format.space_after = Pt(3)
    code_style.paragraph_format.line_spacing = 1.0

    return doc


def add_code(doc, text):
    """Add a code block."""
    for line in text.strip().split("\n"):
        doc.add_paragraph(line, style="Code")


def add_equation(doc, text, label=None):
    """Add a centered equation paragraph."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.name = "Cambria Math"
    run.font.size = Pt(11)
    if label:
        run = p.add_run(f"    ({label})")
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
    return p


def build_report():
    doc = Document()
    doc = setup_styles(doc)

    # Set margins
    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    # =========================================================================
    # TITLE
    # =========================================================================
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("tpacmon: Technical Report")
    run.bold = True
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor(0x1A, 0x3C, 0x6E)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(
        "A Temporal Bayesian Latent Factor Model with Gaussian Process Priors,\n"
        "Gene-Set Informed Horseshoe, and Dual-Level Covariate Regression"
    )
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0x60, 0x60, 0x60)

    date_p = doc.add_paragraph()
    date_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = date_p.add_run(f"Generated: {datetime.date.today().isoformat()}")
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    doc.add_page_break()

    # =========================================================================
    # TABLE OF CONTENTS placeholder
    # =========================================================================
    doc.add_heading("Table of Contents", level=1)
    doc.add_paragraph(
        "1. Overview and Motivation\n"
        "2. The Generative Model\n"
        "    2.1 Factor Loadings: The Regularized Horseshoe\n"
        "    2.2 Temporal Latent Factors: Gaussian Process Priors\n"
        "    2.3 The Smoothness Parameter zeta\n"
        "    2.4 Dual-Level Covariate Regression\n"
        "    2.5 Observation Likelihood\n"
        "3. Variational Inference: How Training Works\n"
        "    3.1 The Evidence Lower Bound (ELBO)\n"
        "    3.2 The Variational Guide (Approximate Posterior)\n"
        "    3.3 Stochastic Variational Inference (SVI)\n"
        "    3.4 The Reparameterization Trick\n"
        "    3.5 What Happens in Each Training Step\n"
        "4. Implementation Details\n"
        "    4.1 Data Representation and Patient Grouping\n"
        "    4.2 Pyro Plates and Tensor Shapes\n"
        "    4.3 GP Reparameterization and Vectorization\n"
        "    4.4 View Scaling\n"
        "    4.5 Optimization, KL Annealing, and Early Stopping\n"
        "5. Post-Training: Extracting Results\n"
        "    5.1 Factor Scores\n"
        "    5.2 Loadings and Sparsity\n"
        "    5.3 GP Prediction at New Timepoints\n"
        "6. Relationship to MuVI, PACMon, MOFA2, and mofaflex"
    )

    doc.add_page_break()

    # =========================================================================
    # 1. OVERVIEW
    # =========================================================================
    doc.add_heading("1. Overview and Motivation", level=1)

    doc.add_paragraph(
        "tpacmon (Temporal PACMON) is a Bayesian latent factor model designed for "
        "longitudinal multi-omics data. It decomposes high-dimensional observations "
        "into a small number of latent factors, each capturing a pattern of "
        "co-variation across features. What distinguishes tpacmon from its "
        "predecessors (MuVI, PACMon, MOFA2) is the integration of three capabilities "
        "into a single coherent generative model:"
    )

    doc.add_paragraph(
        "1. Gaussian Process priors on latent factors, enabling smooth temporal "
        "trajectories with learned dynamics per factor.", style="List Number"
    )
    doc.add_paragraph(
        "2. Gene-set informed regularized horseshoe priors on loadings, guiding "
        "which features each factor should explain based on biological prior knowledge.", style="List Number"
    )
    doc.add_paragraph(
        "3. Dual-level covariate regression: covariates (e.g., sex, treatment) "
        "can shift both the factor trajectory (via gamma) and the feature-level "
        "observations directly (via beta).", style="List Number"
    )

    doc.add_paragraph(
        "The model is trained using Stochastic Variational Inference (SVI) in Pyro. "
        "This section-by-section report explains what each component does, how "
        "they fit together during training, and what the optimizer is actually "
        "learning at each step."
    )

    # =========================================================================
    # 2. GENERATIVE MODEL
    # =========================================================================
    doc.add_heading("2. The Generative Model", level=1)

    doc.add_paragraph(
        "A generative model describes how the observed data could have been "
        "generated from latent (hidden) variables. Think of it as a recipe: "
        '"If you knew the true factor scores, loadings, and noise levels, '
        'here is how you would produce the data." During training, we work '
        "backwards: given the observed data, we infer what the latent variables "
        "most likely were."
    )

    doc.add_paragraph(
        "The core equation is, for each observed visit (patient p, time t, view m):"
    )
    add_equation(doc,
        "y_{p,t,m} = z_{p,t} W_m + x_p beta_m + noise",
        label="1"
    )
    doc.add_paragraph(
        "where z_{p,t} is a K-dimensional factor score vector, W_m is the "
        "(K x D_m) loading matrix for view m, x_p is the patient's covariate "
        "vector, and beta_m captures direct covariate effects on features."
    )

    doc.add_paragraph(
        "The question is: where do z, W, and the other parameters come from? "
        "This is where the prior distributions come in. Each prior encodes our "
        "beliefs about the structure of the latent variables before seeing data."
    )

    # --- 2.1 Horseshoe ---
    doc.add_heading("2.1 Factor Loadings: The Regularized Horseshoe", level=2)

    doc.add_paragraph(
        "The loading matrix W_m determines which features are associated with "
        "each factor. In biology, we expect most gene-factor associations to be "
        "zero (a factor about glycolysis should not load on random immune genes). "
        "The regularized horseshoe prior enforces this sparsity while allowing "
        "truly active features to have large loadings."
    )

    doc.add_paragraph("For each sparse (gene-set informed) factor k, feature j in view m:")

    add_equation(doc, "local_scale_{k,j,m}  ~  HalfCauchy(1)")
    add_equation(doc, "factor_scale_{k,m}    ~  HalfCauchy(1)")
    add_equation(doc, "view_scale_m          ~  HalfCauchy(1)")
    add_equation(doc, "c_aux_{k,j,m}         ~  InverseGamma(0.5, 0.5)")

    doc.add_paragraph(
        "These four scale parameters combine into the effective scale for each loading:"
    )
    add_equation(doc,
        "c = sqrt(c_aux) * prior_scale_{k,j,m}",
        label="2"
    )
    add_equation(doc,
        "tau = local * factor * view    (raw combined scale)",
        label="3"
    )
    add_equation(doc,
        "w_scale = (global * c * tau) / sqrt(c^2 + tau^2)",
        label="4"
    )
    add_equation(doc,
        "w_{k,j,m}  ~  Normal(0, w_scale)",
        label="5"
    )

    doc.add_paragraph(
        "How this achieves sparsity: The horseshoe prior has a sharp spike near "
        "zero (from the HalfCauchy local scales) and heavy tails (allowing large "
        "values). Most loadings get squeezed to near-zero; a few escape with large "
        "magnitudes. The regularization (the c term in Eq. 4) prevents the heavy "
        "tails from producing unrealistically large values by applying a soft "
        "upper bound."
    )

    doc.add_paragraph(
        "How gene-set priors work: The prior_scale_{k,j,m} term encodes biological "
        "knowledge. If gene j is in the gene set for factor k, its prior_scale is "
        "close to 1.0, giving it a wide prior (easy to become active). If gene j "
        "is NOT in the set, its prior_scale is very small (e.g., 0.01), making it "
        "much harder for that gene to escape the spike. Concretely:"
    )
    add_equation(doc,
        "prior_scale = clip(mask + (1 - confidence), 1e-8, 1.0)"
    )
    doc.add_paragraph(
        "where mask is 1 for genes in the set and 0 otherwise. Confidence maps as: "
        "low=0.99, med=0.995, high=0.999. With low confidence (0.99), non-prior "
        "genes get a prior_scale of 0.01 -- small but not impossible to overcome."
    )

    doc.add_paragraph(
        "Dense (uninformed) factors skip the horseshoe entirely. Their loadings "
        "use a simple Normal(0, 1) prior, allowing them to freely capture patterns "
        "not covered by the gene-set factors."
    )

    # --- 2.2 GP ---
    doc.add_heading("2.2 Temporal Latent Factors: Gaussian Process Priors", level=2)

    doc.add_paragraph(
        "In MuVI and PACMon, factor scores z are drawn independently for each "
        "sample: z_{p,k} ~ Normal(0, 1). This means the model has no concept of "
        "temporal ordering -- a patient's factor score at day 1 is unrelated to "
        "their score at day 7."
    )

    doc.add_paragraph(
        "tpacmon replaces this with a Gaussian Process (GP) prior. A GP defines a "
        "distribution over functions, so instead of drawing independent scores, "
        "we draw entire trajectories. The key property: nearby timepoints have "
        "correlated scores, producing smooth curves."
    )

    doc.add_paragraph("For each patient p and factor k:")
    add_equation(doc,
        "f_{p,k}(t)  ~  GP(mu_{p,k}, K_k(t, t'))",
        label="6"
    )

    doc.add_paragraph(
        "This means that the vector of factor scores across all observed timepoints "
        "for patient p, factor k follows a multivariate normal:"
    )
    add_equation(doc,
        "f_{p,k}  ~  MVN(mu_{p,k}, K_k)",
        label="7"
    )

    doc.add_paragraph(
        "where K_k is the kernel (covariance) matrix. Each entry K_k[t_i, t_j] "
        "encodes how correlated the factor score should be between timepoints t_i "
        "and t_j."
    )

    doc.add_heading("Kernel functions", level=3)

    doc.add_paragraph(
        "The default kernel is Matern 3/2, which produces trajectories that are "
        "once-differentiable (smooth but not overly so -- appropriate for "
        "biological dynamics):"
    )
    add_equation(doc,
        "k(r) = a^2 (1 + sqrt(3) r/l) exp(-sqrt(3) r/l)",
        label="8"
    )
    doc.add_paragraph(
        "where r = |t_i - t_j| is the temporal distance, l is the lengthscale, "
        "and a is the amplitude."
    )

    doc.add_paragraph(
        "The lengthscale l controls how quickly correlations decay with temporal "
        "distance. A short lengthscale (l=1) means scores change rapidly over time; "
        "a long lengthscale (l=10) means slow, gradual changes. Crucially, tpacmon "
        "learns a separate lengthscale per factor, so it can discover that some "
        "biological processes change rapidly while others evolve slowly."
    )

    doc.add_paragraph(
        "The amplitude a controls the overall variance of the GP. Both l and a "
        "are given LogNormal priors and learned during training. The amplitude "
        "plays a dual role: it scales the GP covariance (K = a^2 * k_base) and, "
        "via the zeta mixing formula (Eq. 11), also scales the i.i.d. branch to "
        "maintain consistent factor magnitudes across the GP/i.i.d. spectrum."
    )

    doc.add_paragraph(
        "Note: Bayesian factor models have an inherent scale indeterminacy between "
        "factor scores z and loadings W (the product z*c @ W/c gives the same "
        "reconstruction for any scalar c). In tpacmon, the learned amplitude "
        "partially absorbs this ambiguity. While the factor directions and loading "
        "sparsity patterns are reliably recovered, the absolute magnitude of factor "
        "scores should be interpreted relative to the loadings, not in isolation."
    )

    # --- 2.3 Zeta ---
    doc.add_heading("2.3 The Smoothness Parameter zeta", level=2)

    doc.add_paragraph(
        "Not every factor needs temporal structure. Some patterns may be stable "
        "within a patient (e.g., genetic background), making a GP unnecessary. "
        "The smoothness parameter zeta_k, inspired by mofaflex, lets each factor "
        "decide for itself:"
    )
    add_equation(doc, "zeta_k  ~  Beta(1, 1)    (uniform on [0, 1])", label="9")
    add_equation(doc, "eta_{p,k}  ~  Normal(0, 1)    (i.i.d. component)", label="10")
    add_equation(doc,
        "z_{p,k} = sqrt(1 - zeta_k) * f_{p,k} + sqrt(zeta_k) * a_k * eta_{p,k}",
        label="11"
    )

    doc.add_paragraph(
        "When zeta_k approaches 0: the factor is fully temporal -- z equals the GP "
        "draw f, giving smooth trajectories."
    )
    doc.add_paragraph(
        "When zeta_k approaches 1: the factor is fully i.i.d. -- z equals the "
        "amplitude-scaled noise a_k * eta, and the model reverts to standard "
        "PACMon/MuVI behavior for that factor."
    )
    doc.add_paragraph(
        "The sqrt weighting ensures that the total variance of z remains constant "
        "regardless of zeta. The amplitude scaling on the i.i.d. branch is critical: "
        "the GP branch f_{p,k} has marginal variance proportional to a_k^2 (from the "
        "kernel's Cholesky decomposition), so the i.i.d. branch must be scaled by a_k "
        "to keep both branches at the same magnitude. Without this, the model cannot "
        "meaningfully interpolate between GP and i.i.d. regimes."
    )
    doc.add_paragraph(
        "During training, the model learns zeta_k per factor, automatically "
        "discovering which factors benefit from temporal smoothness and which do not."
    )

    # --- 2.4 Covariates ---
    doc.add_heading("2.4 Dual-Level Covariate Regression", level=2)

    doc.add_paragraph(
        "Covariates (e.g., sex, treatment group, age) can influence both the "
        "latent factor trajectories and the observed features directly. tpacmon "
        "models both levels:"
    )

    doc.add_heading("Factor-level: gamma", level=3)
    add_equation(doc,
        "gamma_{c,k}  ~  Normal(0, 1)",
        label="12"
    )
    add_equation(doc,
        "mu_{p,k} = x_p @ gamma[:, k]    (GP mean for patient p, factor k)",
        label="13"
    )
    doc.add_paragraph(
        "This shifts the entire GP trajectory up or down based on patient "
        "covariates. For example, if gamma for sex is 0.5 for factor k, male "
        "patients (x=1) will have their factor-k trajectory shifted up by 0.5 "
        "relative to female patients (x=0). The temporal shape (controlled by the "
        "kernel) remains the same -- only the baseline level changes."
    )

    doc.add_heading("Feature-level: beta", level=3)
    add_equation(doc,
        "beta_{c,j,m}  ~  Normal(0, 1)",
        label="14"
    )
    doc.add_paragraph(
        "Beta captures direct covariate effects on individual features, bypassing "
        "the latent factors entirely. This accounts for covariate-driven variation "
        "that doesn't align with any factor's loading pattern. The full observation "
        "model becomes:"
    )
    add_equation(doc,
        "y_{p,t,m} = z_{p,t} W_m + x_p beta_m + epsilon",
        label="15"
    )

    # --- 2.5 Likelihood ---
    doc.add_heading("2.5 Observation Likelihood", level=2)

    doc.add_paragraph(
        "For Gaussian views (the default), each observed feature value follows:"
    )
    add_equation(doc,
        "y_{p,t,j,m}  ~  Normal(z_{p,t} w_{:,j,m} + x_p beta_{:,j,m},  sigma_{j,m})",
        label="16"
    )
    doc.add_paragraph(
        "where sigma_{j,m} ~ LogNormal(0, 1) is a per-feature noise standard "
        "deviation. Missing observations (unobserved timepoints, NaN features) "
        "are masked out and do not contribute to the likelihood."
    )

    # =========================================================================
    # 3. VARIATIONAL INFERENCE
    # =========================================================================
    doc.add_heading("3. Variational Inference: How Training Works", level=1)

    doc.add_paragraph(
        "The generative model defines a joint distribution p(y, z, W, ...) over "
        "data and latent variables. We want the posterior p(z, W, ... | y) -- the "
        "distribution of latent variables given the observed data. This posterior "
        "is analytically intractable (the integral over all latent variables has "
        "no closed form), so we use variational inference to approximate it."
    )

    # --- 3.1 ELBO ---
    doc.add_heading("3.1 The Evidence Lower Bound (ELBO)", level=2)

    doc.add_paragraph(
        "Variational inference frames posterior inference as an optimization "
        "problem. We posit a family of approximate distributions q(z, W, ...; phi) "
        "parameterized by variational parameters phi, and find the phi that makes "
        "q as close as possible to the true posterior. The objective is the ELBO:"
    )
    add_equation(doc,
        "ELBO(phi) = E_q[ log p(y, theta) - log q(theta; phi) ]",
        label="17"
    )
    doc.add_paragraph(
        "where theta denotes all latent variables collectively. This decomposes as:"
    )
    add_equation(doc,
        "ELBO = E_q[ log p(y | theta) ] - KL( q(theta; phi) || p(theta) )",
        label="18"
    )

    doc.add_paragraph(
        "The first term rewards configurations where the latent variables explain "
        "the data well (high likelihood). The second term penalizes the approximate "
        "posterior for deviating from the prior. Maximizing ELBO is equivalent to "
        "minimizing the KL divergence between q and the true posterior."
    )

    doc.add_paragraph(
        "In practice, Pyro minimizes -ELBO (a loss). Each training epoch computes "
        "a stochastic estimate of the ELBO by:"
    )
    doc.add_paragraph(
        "1. Sampling latent variables from the current guide: theta ~ q(.; phi)", style="List Number"
    )
    doc.add_paragraph(
        "2. Evaluating the model's log-probability at those values: log p(y, theta)", style="List Number"
    )
    doc.add_paragraph(
        "3. Evaluating the guide's log-probability: log q(theta; phi)", style="List Number"
    )
    doc.add_paragraph(
        "4. Computing the difference and taking gradients w.r.t. phi", style="List Number"
    )

    # --- 3.2 Guide ---
    doc.add_heading("3.2 The Variational Guide (Approximate Posterior)", level=2)

    doc.add_paragraph(
        "The guide is the approximate posterior q. For each latent variable in the "
        "model, the guide defines a corresponding variational distribution. In "
        "tpacmon, the guide uses:"
    )

    doc.add_paragraph(
        "Diagonal Normal for: w (loadings), beta (feature-level covariates), "
        "gamma (factor-level covariates), f_eps (GP reparameterization), "
        "eta (i.i.d. component). Each has a learnable mean (loc) and standard "
        "deviation (scale)."
    )
    doc.add_paragraph(
        "Diagonal LogNormal for: lengthscale, amplitude, sigma (noise), "
        "view_scale, factor_scale, local_scale, c_aux. These are positive-valued "
        "variables."
    )
    doc.add_paragraph(
        "Beta distribution for: zeta (smoothness). Parameterized by two learnable "
        "positive concentrations alpha, beta."
    )

    doc.add_paragraph(
        '"Diagonal" means the guide assumes independence between all parameters '
        "(mean-field approximation). For the GP factor scores specifically, early "
        "versions of tpacmon used a full Cholesky-parameterized MVN (capturing "
        "temporal correlations in the posterior), but this was replaced with a "
        "diagonal Normal for memory efficiency. The temporal correlations are "
        "instead captured through the reparameterization trick (Section 3.4)."
    )

    doc.add_paragraph(
        "Each variational parameter is stored as a PyroParam with appropriate "
        "constraints (e.g., positive for scales). The total number of variational "
        "parameters is proportional to the number of model parameters -- roughly "
        "2x (a loc and a scale for each)."
    )

    # --- 3.3 SVI ---
    doc.add_heading("3.3 Stochastic Variational Inference (SVI)", level=2)

    doc.add_paragraph(
        "SVI combines the ELBO objective with stochastic gradient descent. Pyro's "
        "SVI class orchestrates this:"
    )
    add_code(doc, """\
svi = pyro.infer.SVI(
    model=temporal_model,     # the generative model (forward pass)
    guide=temporal_guide,     # the approximate posterior
    optim=ClippedAdam(...),   # optimizer for variational params
    loss=Trace_ELBO(),        # ELBO estimator
)
loss = svi.step(obs, obs_mask, patient_idx, time_idx, covs)""")

    doc.add_paragraph(
        "Each call to svi.step() performs one complete cycle:"
    )
    doc.add_paragraph(
        "1. Run the guide forward: sample all latent variables from the current "
        "approximate posterior q(theta; phi). This produces concrete values for "
        "every z, W, sigma, lengthscale, etc.", style="List Number"
    )
    doc.add_paragraph(
        "2. Run the model forward with these sampled values: compute the joint "
        "log-probability log p(y, theta) by evaluating each prior and the "
        "likelihood.", style="List Number"
    )
    doc.add_paragraph(
        "3. Compute the ELBO estimate: subtract the guide's log-probability from "
        "the model's log-probability.", style="List Number"
    )
    doc.add_paragraph(
        "4. Backpropagate: compute gradients of -ELBO w.r.t. all variational "
        "parameters phi (the loc and scale parameters in the guide).", style="List Number"
    )
    doc.add_paragraph(
        "5. Update parameters: the optimizer (ClippedAdam) takes a gradient step "
        "to improve phi.", style="List Number"
    )

    doc.add_paragraph(
        "This is repeated for n_epochs iterations. The ELBO should increase "
        "(loss decrease) over training, indicating that the approximate posterior "
        "is getting closer to the true posterior."
    )

    # --- 3.4 Reparam ---
    doc.add_heading("3.4 The Reparameterization Trick", level=2)

    doc.add_paragraph(
        "A critical implementation detail for the GP factors. The generative model "
        "says f_{p,k} ~ MVN(mu, K_k), but sampling directly from an MVN inside "
        "the model would make it hard to backpropagate gradients through the "
        "kernel parameters (lengthscale, amplitude)."
    )

    doc.add_paragraph(
        "Instead, tpacmon uses the reparameterization trick:"
    )
    add_equation(doc, "eps  ~  Normal(0, I)    (standard normal, no learnable params)")
    add_equation(doc, "L = cholesky(K_k)       (lower triangular decomposition of kernel)")
    add_equation(doc, "f = mu + L @ eps         (transform to GP sample)", label="19")

    doc.add_paragraph(
        "This is mathematically equivalent to sampling from MVN(mu, K), but the "
        "randomness is isolated in eps (which has no parameters), while L and mu "
        "are deterministic functions of the learnable parameters. Gradients can "
        "now flow through L (and thus through lengthscale and amplitude) via "
        "standard backpropagation."
    )

    doc.add_paragraph(
        "In the code, the model samples eps as a standard Normal and the guide "
        "provides a diagonal Normal approximation to eps. The actual GP sample "
        "is reconstructed outside the sample statement:"
    )
    add_code(doc, """\
# Model: sample eps ~ Normal(0, I), shape (n_patients_in_group, K, T)
f_eps = pyro.sample("f_g0", Normal(zeros, ones).to_event(2))

# Reconstruct: f = mu + L @ eps  (batched over patients via einsum)
L_all = torch.linalg.cholesky(K_all)  # (K, T, T)
f_group = gp_means + einsum("kij,nkj->nki", L_all, f_eps)""")

    doc.add_paragraph(
        "The einsum operation is the batched matrix multiplication: for each "
        "factor k, it multiplies L_k (T x T) by each patient's eps_k (T,), "
        "producing the GP sample for all patients and factors simultaneously."
    )

    # --- 3.5 Training step ---
    doc.add_heading("3.5 What Happens in Each Training Step", level=2)

    doc.add_paragraph(
        "Putting it all together, here is the complete flow of a single training "
        "step, tracing what happens to each component:"
    )

    doc.add_paragraph(
        "Guide forward pass (sampling from q):", style="List Bullet"
    )
    doc.add_paragraph(
        "  - Sample horseshoe scales: local_scale, factor_scale, view_scale, c_aux "
        "from their LogNormal variational distributions"
    )
    doc.add_paragraph(
        "  - Sample loadings w from Normal(loc_w, scale_w)"
    )
    doc.add_paragraph(
        "  - Sample GP hyperparameters: lengthscale, amplitude from LogNormal; "
        "zeta from Beta"
    )
    doc.add_paragraph(
        "  - Sample gamma (covariate-factor effects) from Normal"
    )
    doc.add_paragraph(
        "  - Sample f_eps (GP reparameterization noise) from Normal(loc_eps, scale_eps) "
        "for each patient group"
    )
    doc.add_paragraph(
        "  - Sample eta (i.i.d. component) from Normal(loc_eta, scale_eta)"
    )

    doc.add_paragraph(
        "Model forward pass (evaluating log p):", style="List Bullet"
    )
    doc.add_paragraph(
        "  - Evaluate horseshoe prior log-densities on the sampled scales"
    )
    doc.add_paragraph(
        "  - Compute effective w_scale (Eq. 4) and evaluate Normal prior on w"
    )
    doc.add_paragraph(
        "  - Build kernel matrices K_k from sampled lengthscales and amplitudes"
    )
    doc.add_paragraph(
        "  - Compute GP means from gamma and covariates (Eq. 13)"
    )
    doc.add_paragraph(
        "  - Evaluate Normal(0, I) prior on f_eps (the reparameterized GP noise)"
    )
    doc.add_paragraph(
        "  - Reconstruct f = mu + L @ eps (Eq. 19)"
    )
    doc.add_paragraph(
        "  - Mix: z = sqrt(1-zeta) * f + sqrt(zeta) * amplitude * eta (Eq. 11)"
    )
    doc.add_paragraph(
        "  - Compute y_pred = z @ W + x @ beta (Eq. 15)"
    )
    doc.add_paragraph(
        "  - Evaluate observation likelihood: sum of log Normal(y_obs; y_pred, sigma) "
        "over all observed entries"
    )

    doc.add_paragraph(
        "ELBO computation and gradient update:", style="List Bullet"
    )
    doc.add_paragraph(
        "  - ELBO = sum(model log-probs) - sum(guide log-probs)"
    )
    doc.add_paragraph(
        "  - Compute gradients of -ELBO w.r.t. all variational parameters"
    )
    doc.add_paragraph(
        "  - Clip gradients (max norm 10.0) to prevent instability from horseshoe "
        "and GP interactions"
    )
    doc.add_paragraph(
        "  - ClippedAdam updates all variational parameters (locs and scales)"
    )

    doc.add_paragraph(
        "Over many epochs, this process drives the variational parameters toward "
        "values where: (a) the approximate posterior explains the data well "
        "(high likelihood term), (b) the inferred latent variables are plausible "
        "under the priors (low KL term), and (c) the GP hyperparameters (lengthscale, "
        "amplitude, zeta) reflect the actual temporal structure in the data."
    )

    # =========================================================================
    # 4. IMPLEMENTATION DETAILS
    # =========================================================================
    doc.add_heading("4. Implementation Details", level=1)

    # --- 4.1 Data ---
    doc.add_heading("4.1 Data Representation and Patient Grouping", level=2)

    doc.add_paragraph(
        "Longitudinal multi-omics data has shape (P patients, T timepoints, "
        "D features) per view, but not all patients are observed at all timepoints. "
        "tpacmon handles this via a common time grid with per-patient boolean masks."
    )

    doc.add_paragraph(
        "For the SVI training loop, observations are flattened to (N_obs, D_total) "
        "where N_obs = sum of observed visits across all patients, and D_total is "
        "the concatenation of all views. Each row is one patient-visit, tracked by "
        "patient_idx and time_idx arrays."
    )

    doc.add_paragraph(
        "For the GP sampling, patients are grouped by their observation mask "
        "pattern. All patients with identical mask patterns share the same kernel "
        "matrix K_k (since they observe the same timepoints), enabling batched "
        "computation. With 4 timepoints, there are at most 2^4 = 16 possible "
        "patterns; in practice (especially when all patients share the same "
        "schedule), there is often just one group."
    )

    # --- 4.2 Plates ---
    doc.add_heading("4.2 Pyro Plates and Tensor Shapes", level=2)

    doc.add_paragraph(
        "Pyro plates declare dimensions that are conditionally independent, "
        "enabling the ELBO to be computed correctly. The key plates in tpacmon are:"
    )

    table = doc.add_table(rows=7, cols=3)
    table.style = "Light Grid Accent 1"
    headers = ["Plate name", "Size", "Purpose"]
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    data_rows = [
        ("factor_left_sparse", "K_sparse", "Informed factor dimension for loadings"),
        ("factor_left_dense", "K_dense", "Uninformed factor dimension for loadings"),
        ("feature_m", "D_m", "Feature dimension per view"),
        ("factor", "K", "Factor dimension for GP hyperparameters"),
        ("obs", "N_obs", "Observation (patient-visit) dimension"),
        ("group_g", "N_group", "Patient group dimension for GP samples"),
    ]
    for row_idx, (name, size, purpose) in enumerate(data_rows, 1):
        table.rows[row_idx].cells[0].text = name
        table.rows[row_idx].cells[1].text = size
        table.rows[row_idx].cells[2].text = purpose

    doc.add_paragraph("")
    doc.add_paragraph(
        "The .to_event(2) call on the GP distributions (f_eps, eta) tells Pyro "
        "that the last 2 dimensions (K, T) are event dimensions (not independent), "
        "leaving only the patient group dimension as the batch dimension claimed by "
        "the plate."
    )

    # --- 4.3 GP vectorization ---
    doc.add_heading("4.3 GP Reparameterization and Vectorization", level=2)

    doc.add_paragraph(
        "The naive implementation would loop over P patients and K factors, "
        "issuing P*K separate pyro.sample calls. This is prohibitively slow for "
        "realistic datasets (e.g., 95 patients x 600 factors = 57,000 sample sites)."
    )

    doc.add_paragraph("tpacmon uses two optimizations:")

    doc.add_paragraph(
        "1. Patient grouping: patients sharing the same observation mask are "
        "batched together. The kernel matrix K_k only depends on which timepoints "
        "are observed, so all patients in a group share the same K and L matrices.", style="List Number"
    )
    doc.add_paragraph(
        "2. Factor batching via reparameterization: instead of sampling from "
        "MVN(mu_k, K_k) for each factor separately, we sample a single block "
        "eps ~ Normal(0, I) of shape (n_group, K, T) and reconstruct all factors "
        "simultaneously via f = mu + L @ eps. This reduces 2*K sample sites per "
        "group to just 2 (one for f_eps, one for eta).", style="List Number"
    )

    # --- 4.4 View scaling ---
    doc.add_heading("4.4 View Scaling", level=2)

    doc.add_paragraph(
        "When views have different numbers of features (e.g., 5000 genes vs. 200 "
        "proteins), the view with more features dominates the ELBO simply because "
        "it contributes more likelihood terms. tpacmon balances this by scaling "
        "each view's contribution:"
    )
    add_equation(doc,
        "scale_m = M * (1 / D_m^alpha) / sum_m(1 / D_m^alpha),   alpha = 0.25",
        label="20b"
    )
    doc.add_paragraph(
        "This downweights large views and upweights small ones, ensuring that "
        "all views contribute meaningfully to factor learning. The scaling is "
        "applied via pyro.poutine.scale() around the likelihood terms."
    )

    # --- 4.5 Optimization ---
    doc.add_heading("4.5 Optimization and Early Stopping", level=2)

    doc.add_paragraph(
        "Optimizer: ClippedAdam with exponential learning rate decay. The initial "
        "learning rate (default 0.01) decays by a factor of gamma=0.1 over the "
        "total number of epochs: lr_t = lr_0 * 0.1^(t/T). Gradient clipping at "
        "max_norm=10.0 prevents instability from the heavy-tailed horseshoe and "
        "GP interactions."
    )

    doc.add_heading("KL annealing", level=3)
    doc.add_paragraph(
        "The factor score priors (f_eps and eta, both Normal(0,1)) are wrapped "
        "in a KL weight that ramps linearly from 0.01 to 1.0 over the first 25% "
        "of training (capped at 500 epochs). This prevents posterior collapse -- "
        "a known failure mode in mean-field variational inference where the KL "
        "penalty on hundreds of factor score parameters overwhelms the reconstruction "
        "gradient, causing the guide to stay near the prior (zero mean, small scale) "
        "rather than learning informative factor scores."
    )
    add_equation(doc,
        "kl_weight(t) = min(1.0, 0.01 + 0.99 * t / anneal_epochs)",
        label="20a"
    )
    doc.add_paragraph(
        "During early training, the low KL weight lets factors learn their correct "
        "magnitudes driven purely by reconstruction quality. As the weight increases, "
        "the KL term gradually regularizes the posterior, preventing overfitting."
    )

    doc.add_heading("Early stopping with ELBO smoothing", level=3)
    doc.add_paragraph(
        "Training halts when the smoothed ELBO improvement falls below a tolerance "
        "(1e-5) for a patience window (default 100 epochs), with a minimum of 500 "
        "epochs. The smoothing uses a moving average (window=50 epochs) of the ELBO "
        "before comparing to the best seen value. This is necessary because the SVI "
        "ELBO is inherently noisy -- it is a stochastic estimate from a single sample "
        "per epoch. Without smoothing, epoch-to-epoch fluctuations of hundreds of "
        "units trigger premature stopping even when the trend is still improving."
    )

    # =========================================================================
    # 5. POST-TRAINING
    # =========================================================================
    doc.add_heading("5. Post-Training: Extracting Results", level=1)

    doc.add_heading("5.1 Factor Scores", level=2)
    doc.add_paragraph(
        "After training, factor scores are extracted from the guide's variational "
        "means. The guide stores z_mean as a (P, K, T_max) tensor. The method "
        "get_factors() maps these back to the (P, T, K) array aligned with the "
        "original time grid, filling NaN for unobserved timepoints."
    )

    doc.add_heading("5.2 Loadings and Sparsity", level=2)
    doc.add_paragraph(
        "Loadings are the variational mode (mean for Normal-distributed sites) "
        "of the w parameters. For LogNormal-distributed sites (like scales), the "
        "mode is exp(mu - sigma^2). The horseshoe prior will have driven most "
        "loadings to near-zero during training; the nonzero loadings identify "
        "the active features per factor."
    )

    doc.add_heading("5.3 GP Prediction at New Timepoints", level=2)
    doc.add_paragraph(
        "Given a trained model, we can predict factor scores at unobserved "
        "timepoints using GP conditional distributions. For new times t_new:"
    )
    add_equation(doc,
        "f(t_new) | f(t_obs)  ~  Normal(K_new_obs K_obs^{-1} f_obs,  "
        "K_new - K_new_obs K_obs^{-1} K_new_obs^T)",
        label="21"
    )
    doc.add_paragraph(
        "This gives both a mean prediction and uncertainty (variance) at each "
        "new timepoint. The predict() method implements this using the learned "
        "kernel parameters and the variational mean of the factor scores at "
        "observed times."
    )

    # =========================================================================
    # 6. RELATIONSHIP TO OTHER MODELS
    # =========================================================================
    doc.add_heading("6. Relationship to MuVI, PACMon, MOFA2, and mofaflex", level=1)

    doc.add_paragraph(
        "tpacmon sits in a lineage of Bayesian multi-omics factor models, each "
        "adding capabilities:"
    )

    table2 = doc.add_table(rows=6, cols=2)
    table2.style = "Light Grid Accent 1"
    table2.rows[0].cells[0].text = "Model"
    table2.rows[0].cells[1].text = "What it adds"
    rows2 = [
        ("MOFA2", "Multi-view factor model with ARD sparsity. No gene-set priors, no covariates, no temporal structure."),
        ("MEFISTO", "Adds GP temporal structure to MOFA2, but only RBF kernel, grid-search lengthscales, and Kronecker structure forcing shared time grids."),
        ("MuVI", "Adds gene-set informed priors (horseshoe) and covariate regression to MOFA2. No temporal structure."),
        ("PACMon", "Extends MuVI with regularized horseshoe, condition-specific effects, and improved sparsity. No temporal structure."),
        ("tpacmon", "Extends PACMon with GP temporal structure (per-factor kernels, learned lengthscales, zeta smoothness mixing), irregular time support via patient masking, and factor-level covariate-shifted GP means (gamma)."),
    ]
    for i, (model, desc) in enumerate(rows2, 1):
        table2.rows[i].cells[0].text = model
        table2.rows[i].cells[1].text = desc

    doc.add_paragraph("")
    doc.add_paragraph(
        "The key insight of tpacmon is that temporal structure (GP), sparsity "
        "structure (horseshoe), and covariate structure (gamma/beta) are not "
        "independent concerns. By placing them in a single generative model and "
        "training jointly via SVI, the model can discover interactions: for example, "
        "a gene-set factor whose temporal dynamics differ by treatment group. This "
        "is not possible when GP and horseshoe operate independently (as in mofaflex) "
        "or when temporal structure is absent entirely (as in PACMon/MuVI)."
    )

    # =========================================================================
    # SAVE
    # =========================================================================
    outpath = "/Users/sebastian/Library/CloudStorage/OneDrive-Persönlich/Forschung/Temporal_muvi_flex/tpacmon_technical_report.docx"
    doc.save(outpath)
    print(f"Report saved to: {outpath}")


if __name__ == "__main__":
    build_report()
